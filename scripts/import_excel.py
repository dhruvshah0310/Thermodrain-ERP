#!/usr/bin/env python3
"""
One-time migration: loads the legacy Purchase Excel (INDENT_REPORT) into the
Thermodrain ERP SQLite database, so the Purchase team starts with continuity
instead of an empty system.

The Excel's MATERIAL sheet is a flat, line-item-per-row export (one row per
material per purchase-order-with-receipt) — it does not separately record the
Indent -> Sanction Indent -> RFQ -> Quotation -> PO -> GRN steps the way the
live TSS software does. This script reconstructs a reasonable document trail
from it:
  - one purchase_indent + sanction_indent per distinct TSS Indent No.
  - one purchase_order per distinct Purchase Order No. (its real legacy
    number is kept as doc_no, so old paperwork stays searchable)
  - one grn per distinct GRN No. (GRN columns map almost 1:1 to our schema,
    including the already-computed landed rate)

Run with: python3 scripts/import_excel.py <path-to-xlsx> [--db path/to/erp.db]
"""
import argparse
import re
import sqlite3
import sys
from collections import defaultdict
from datetime import datetime, date
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Missing dependency: pip install openpyxl")

REPO_ROOT = Path(__file__).resolve().parent.parent
SCHEMA_PATH = REPO_ROOT / "server" / "src" / "db" / "schema.sql"
DEFAULT_DB_PATH = REPO_ROOT / "server" / "data" / "erp.db"

LEGACY_DEPARTMENT = "STORE"  # MATERIAL sheet has no department column; the video shows STORE as the default indent department.


def to_iso_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value)


def slugify_code(name: str, seq: int) -> str:
    base = re.sub(r"[^A-Z0-9]", "", (name or "MAT").upper())[:10] or "MAT"
    return f"{base}-{seq:04d}"


def get_or_create_department(cur, cache, name):
    if name in cache:
        return cache[name]
    cur.execute("SELECT id FROM departments WHERE name = ?", (name,))
    row = cur.fetchone()
    if row:
        cache[name] = row[0]
        return row[0]
    cur.execute("INSERT INTO departments (name, code) VALUES (?, ?)", (name, name[:4].upper()))
    cache[name] = cur.lastrowid
    return cur.lastrowid


def get_or_create_unit(cur, cache, name):
    name = (name or "NOS").strip() or "NOS"
    if name in cache:
        return cache[name]
    cur.execute("SELECT id FROM units WHERE name = ?", (name,))
    row = cur.fetchone()
    if row:
        cache[name] = row[0]
        return row[0]
    cur.execute("INSERT INTO units (name) VALUES (?)", (name,))
    cache[name] = cur.lastrowid
    return cur.lastrowid


def get_or_create_vendor(cur, cache, name):
    name = (name or "").strip()
    if not name:
        return None
    if name in cache:
        return cache[name]
    cur.execute("SELECT id FROM accounts WHERE kind = 'vendor' AND name = ?", (name,))
    row = cur.fetchone()
    if row:
        cache[name] = row[0]
        return row[0]
    cur.execute("INSERT INTO accounts (kind, name) VALUES ('vendor', ?)", (name,))
    cache[name] = cur.lastrowid
    return cur.lastrowid


def get_or_create_material(cur, cache, seq_holder, name, material_type, group_head, sub_group, unit_id):
    name = (name or "").strip()
    if not name:
        return None
    key = name
    if key in cache:
        return cache[key]
    cur.execute("SELECT id FROM materials WHERE name = ?", (name,))
    row = cur.fetchone()
    if row:
        cache[key] = row[0]
        return row[0]
    seq_holder[0] += 1
    code = slugify_code(sub_group or material_type or name, seq_holder[0])
    cur.execute(
        """INSERT INTO materials (code, name, unit_id, material_type, group_head, sub_group)
           VALUES (?, ?, ?, ?, ?, ?)""",
        (code, name, unit_id, material_type, group_head, sub_group),
    )
    cache[key] = cur.lastrowid
    return cur.lastrowid


def fiscal_year_label(d: date) -> str:
    start_year = d.year if d.month >= 4 else d.year - 1
    return f"{str(start_year)[-2:]}-{str(start_year + 1)[-2:]}"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx_path")
    parser.add_argument("--db", default=str(DEFAULT_DB_PATH))
    args = parser.parse_args()

    db_path = Path(args.db)
    db_path.parent.mkdir(parents=True, exist_ok=True)

    conn = sqlite3.connect(db_path)
    conn.executescript(SCHEMA_PATH.read_text())
    cur = conn.cursor()
    cur.execute("PRAGMA foreign_keys = ON")

    wb = openpyxl.load_workbook(args.xlsx_path, data_only=True)
    ws = wb["MATERIAL"]
    headers = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
    # The sheet repeats PARTY NAME / MATERIAL TYPE / UNIT / RECEIVE-NOT / SUB GROUP
    # a second time after a blank spacer column — that second block is an
    # unrelated dropdown-validation reference list, not per-row data. Keep
    # only the FIRST occurrence of each header name (the real transaction
    # columns) so duplicates don't silently overwrite them.
    col = {}
    for i, h in enumerate(headers):
        if h and h not in col:
            col[h] = i + 1

    def cell(row, name):
        c = col.get(name)
        return ws.cell(row=row, column=c).value if c else None

    dept_cache, unit_cache, vendor_cache, material_cache = {}, {}, {}, {}
    material_seq = [0]
    taxonomy_seen = set()

    indent_no_to_row_ids = defaultdict(list)
    po_no_to_row_ids = defaultdict(list)
    grn_no_to_row_ids = defaultdict(list)
    rows = []

    print("Reading MATERIAL sheet...")
    skipped_blank = 0
    skipped_placeholder = 0
    for r in range(3, ws.max_row + 1):
        indent_no = cell(r, "TSS INDENT NO.")
        material_details = cell(r, "MATERIAL DETAILS")
        # Many rows in this export are leftover template rows: a dangling
        # "TP-" / "PO/TSS/.../" placeholder with no material, qty, or amount
        # at all. They aren't real transactions, so importing them would
        # create a single garbage document with thousands of empty lines.
        if not indent_no:
            skipped_blank += 1
            continue
        if not material_details:
            skipped_placeholder += 1
            continue
        row = {
            "row": r,
            "indent_no": str(indent_no).strip(),
            "indent_date": to_iso_date(cell(r, "INDENT DATE")),
            "po_no": cell(r, "PURCHASE ORDER NO."),
            "po_date": to_iso_date(cell(r, "PO DATE")),
            "party_name": cell(r, "PARTY NAME"),
            "material_type": cell(r, "MATERIAL TYPE"),
            "group_head": cell(r, "GROUP HEAD"),
            "sub_group": cell(r, "SUB GROUP"),
            "material_details": cell(r, "MATERIAL DETAILS"),
            "qty": cell(r, "QTY.") or 0,
            "unit": cell(r, "UNIT") or "NOS",
            "rate": cell(r, "RATE") or 0,
            "net_amount": cell(r, "NET AMOUNT") or 0,
            "gst_freight": cell(r, "GST FREIGHT") or 0,
            "gst_percent": (cell(r, "GST %") or 0),
            "gst_amount": cell(r, "GST AMOUNT") or 0,
            "without_gst_freight": cell(r, "WITHOUT GST FREIGHT") or 0,
            "grand_total": cell(r, "GRAND TOTAL") or 0,
            "bill_no": cell(r, "BILL NO."),
            "bill_date": to_iso_date(cell(r, "BILL DATE")),
            "challan_no": cell(r, "CHALLAN NO."),
            "challan_date": to_iso_date(cell(r, "CHALLAN DATE")),
            "grn_no": cell(r, "GRN NO."),
            "grn_date": to_iso_date(cell(r, "GRN DATE")),
            "receive_status": (cell(r, "RECEIVE/ NOT") or "").strip().upper() or None,
            "receive_date": to_iso_date(cell(r, "RECEIVE DATE")),
            "receive_qty": cell(r, "RECEIVE QTY.") or 0,
            "basic_receive_amount": cell(r, "BASIC RECEIVE AMOUNT") or 0,
            "freight_gst_recd": cell(r, "FREIGHT +GST+RECD.BASIC AMOUNT") or 0,
            "landed_rate": cell(r, "LANDED RATE") or 0,
        }
        rows.append(row)
        indent_no_to_row_ids[row["indent_no"]].append(len(rows) - 1)
        # A PO/GRN number ending in "/" (e.g. "PO/TSS/26-27/") is the same
        # kind of unfilled placeholder as the blank/"TP-" indents above — the
        # indent exists but hasn't actually been converted to a PO/received yet.
        po_no = str(row["po_no"]).strip() if row["po_no"] else ""
        if po_no and not po_no.endswith("/"):
            po_no_to_row_ids[po_no].append(len(rows) - 1)
        grn_no = str(row["grn_no"]).strip() if row["grn_no"] else ""
        if grn_no and not grn_no.endswith("/"):
            grn_no_to_row_ids[grn_no].append(len(rows) - 1)

        if row["material_type"]:
            taxonomy_seen.add((row["material_type"], row["group_head"], row["sub_group"]))

    print(f"{len(rows)} line items, {len(indent_no_to_row_ids)} indents, {len(po_no_to_row_ids)} POs, {len(grn_no_to_row_ids)} GRNs")
    print(f"Skipped {skipped_blank} rows with no indent no., {skipped_placeholder} rows with an indent no. but no material (empty template rows)")

    dept_id = get_or_create_department(cur, dept_cache, LEGACY_DEPARTMENT)

    print("Seeding material taxonomy...")
    for mt, gh, sg in taxonomy_seen:
        cur.execute(
            "INSERT OR IGNORE INTO material_taxonomy (material_type, group_head, sub_group) VALUES (?, ?, ?)",
            (mt, gh, sg),
        )

    # ---- Indents + Sanction Indents (one pair per distinct TSS Indent No.) ----
    print("Importing indents + sanction indents...")
    indent_db_id = {}
    sanction_db_id = {}
    for indent_no, row_idxs in indent_no_to_row_ids.items():
        first = rows[row_idxs[0]]
        cur.execute(
            """INSERT OR IGNORE INTO purchase_indents (indent_no, indent_date, department_id, narration, status)
               VALUES (?, ?, ?, 'Imported from legacy Excel', 'sanctioned')""",
            (indent_no, first["indent_date"] or "2023-04-01", dept_id),
        )
        cur.execute("SELECT id FROM purchase_indents WHERE indent_no = ?", (indent_no,))
        pi_id = cur.fetchone()[0]
        indent_db_id[indent_no] = pi_id

        sanction_doc_no = f"SI-LEGACY/{indent_no}"
        cur.execute(
            """INSERT OR IGNORE INTO sanction_indents (doc_no, doc_date, department_id, indent_id, narration, purchase_order_linked)
               VALUES (?, ?, ?, ?, 'Imported from legacy Excel', 1)""",
            (sanction_doc_no, first["indent_date"] or "2023-04-01", dept_id, pi_id),
        )
        cur.execute("SELECT id FROM sanction_indents WHERE doc_no = ?", (sanction_doc_no,))
        si_id = cur.fetchone()[0]
        sanction_db_id[indent_no] = si_id

        for i, idx in enumerate(row_idxs):
            row = rows[idx]
            unit_id = get_or_create_unit(cur, unit_cache, row["unit"])
            mat_id = get_or_create_material(
                cur, material_cache, material_seq, row["material_details"],
                row["material_type"], row["group_head"], row["sub_group"], unit_id,
            )
            cur.execute(
                """INSERT INTO purchase_indent_items (indent_id, sr_no, material_id, unit, req_qty)
                   VALUES (?, ?, ?, ?, ?)""",
                (pi_id, i + 1, mat_id, row["unit"], row["qty"]),
            )
            cur.execute(
                """INSERT INTO sanction_indent_items (sanction_id, sr_no, material_id, unit, qty)
                   VALUES (?, ?, ?, ?, ?)""",
                (si_id, i + 1, mat_id, row["unit"], row["qty"]),
            )
            row["material_id"] = mat_id

    # ---- Purchase Orders (one per distinct PO No., keeping the legacy number) ----
    print("Importing purchase orders...")
    po_db_id = {}
    for po_no, row_idxs in po_no_to_row_ids.items():
        first = rows[row_idxs[0]]
        vendor_id = get_or_create_vendor(cur, vendor_cache, first["party_name"])
        sanction_id = sanction_db_id.get(first["indent_no"])

        total_qty = sum(rows[i]["qty"] for i in row_idxs)
        assessable_amt = sum(rows[i]["net_amount"] for i in row_idxs)
        tax_amt = sum(rows[i]["gst_amount"] for i in row_idxs)
        freight = sum(rows[i]["gst_freight"] + rows[i]["without_gst_freight"] for i in row_idxs)
        total_amt = sum(rows[i]["grand_total"] for i in row_idxs)

        cur.execute(
            """INSERT OR IGNORE INTO purchase_orders
                 (doc_no, doc_date, vendor_id, department_id, sanction_id, si_date, narration,
                  cgst, sgst, tax_amt, freight, total_qty, total_amt, gross_amt, assessable_amt, status)
               VALUES (?, ?, ?, ?, ?, ?, 'Imported from legacy Excel', ?, ?, ?, ?, ?, ?, ?, ?, 'issued')""",
            (
                po_no, first["po_date"] or first["indent_date"] or "2023-04-01", vendor_id, dept_id, sanction_id,
                first["indent_date"], tax_amt / 2, tax_amt / 2, tax_amt, freight, total_qty, total_amt, total_amt, assessable_amt,
            ),
        )
        cur.execute("SELECT id FROM purchase_orders WHERE doc_no = ?", (po_no,))
        po_id = cur.fetchone()[0]
        po_db_id[po_no] = po_id

        for i, idx in enumerate(row_idxs):
            row = rows[idx]
            unit_id = get_or_create_unit(cur, unit_cache, row["unit"])
            mat_id = row.get("material_id") or get_or_create_material(
                cur, material_cache, material_seq, row["material_details"],
                row["material_type"], row["group_head"], row["sub_group"], unit_id,
            )
            cur.execute(
                """INSERT INTO purchase_order_items
                     (po_id, sr_no, material_id, description, unit, qty, rate, amount, received_qty)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (po_id, i + 1, mat_id, row["material_details"], row["unit"], row["qty"], row["rate"], row["net_amount"], 0),
            )
            row["po_item_id"] = cur.lastrowid
            row["po_id"] = po_id

    # ---- GRNs (one per distinct GRN No.) ----
    print("Importing GRNs...")
    for grn_no, row_idxs in grn_no_to_row_ids.items():
        first = rows[row_idxs[0]]
        po_id = po_db_id.get(str(first["po_no"]).strip()) if first["po_no"] else None
        status = first["receive_status"] or "RECEIVED"
        if status not in ("RECEIVED", "HOLD", "CANCEL", "NOT"):
            status = "RECEIVED"
        cur.execute(
            """INSERT OR IGNORE INTO grns (grn_no, grn_date, po_id, bill_no, bill_date, challan_no, challan_date, receive_status, receive_date)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                grn_no, first["grn_date"] or first["po_date"] or "2023-04-01", po_id, first["bill_no"], first["bill_date"],
                first["challan_no"], first["challan_date"], status, first["receive_date"],
            ),
        )
        cur.execute("SELECT id FROM grns WHERE grn_no = ?", (grn_no,))
        grn_row = cur.fetchone()
        if not grn_row:
            continue
        grn_id = grn_row[0]

        for idx in row_idxs:
            row = rows[idx]
            po_item_id = row.get("po_item_id")
            mat_id = row.get("material_id")
            cur.execute(
                """INSERT INTO grn_items
                     (grn_id, po_item_id, material_id, receive_qty, rate, basic_receive_amount, freight_gst_recd_basic_amount, landed_rate)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
                (grn_id, po_item_id, mat_id, row["receive_qty"], row["rate"], row["basic_receive_amount"], row["freight_gst_recd"], row["landed_rate"]),
            )
            if po_item_id and status != "CANCEL":
                cur.execute(
                    "UPDATE purchase_order_items SET received_qty = received_qty + ? WHERE id = ?",
                    (row["receive_qty"], po_item_id),
                )

    # Recompute PO status from received quantities.
    cur.execute("""
        UPDATE purchase_orders SET status = (
          SELECT CASE
            WHEN COALESCE(SUM(received_qty),0) <= 0 THEN 'issued'
            WHEN COALESCE(SUM(received_qty),0) >= COALESCE(SUM(qty),0) THEN 'received'
            ELSE 'partially_received'
          END
          FROM purchase_order_items WHERE po_id = purchase_orders.id
        )
    """)

    conn.commit()

    n_vendors = cur.execute("SELECT COUNT(*) FROM accounts WHERE kind='vendor'").fetchone()[0]
    n_materials = cur.execute("SELECT COUNT(*) FROM materials").fetchone()[0]
    n_indents = cur.execute("SELECT COUNT(*) FROM purchase_indents").fetchone()[0]
    n_pos = cur.execute("SELECT COUNT(*) FROM purchase_orders").fetchone()[0]
    n_grns = cur.execute("SELECT COUNT(*) FROM grns").fetchone()[0]
    print(f"Done. vendors={n_vendors} materials={n_materials} indents={n_indents} purchase_orders={n_pos} grns={n_grns}")
    conn.close()


if __name__ == "__main__":
    main()
